"""M1 - text extraction and normalization (resume file -> clean text T)."""
import re
import unicodedata
from pathlib import Path


def normalize_text(text: str) -> str:
    """Unicode-normalize, fix hyphenated line breaks, collapse whitespace.

    Newlines are kept as single line breaks (resume headers are line-based)."""
    text = unicodedata.normalize("NFKC", text)
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", " ", text)   # control chars
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)                  # de-hyphenate
    text = re.sub(r"[ \t\u00a0]+", " ", text)
    text = re.sub(r"\n\s*\n+", "\n", text)
    return "\n".join(line.strip() for line in text.split("\n")).strip()


def _read_pdf(path: Path) -> str:
    try:
        import pdfplumber
        with pdfplumber.open(str(path)) as pdf:
            pages = []
            for p in pdf.pages:
                text = p.extract_text(layout=True) or p.extract_text() or ""
                pages.append(text)
            return "\n".join(pages)
    except ImportError:
        pass
    except Exception:
        pass

    try:
        from pypdf import PdfReader
        return "\n".join((p.extract_text() or "") for p in PdfReader(str(path)).pages)
    except ImportError:
        raise ImportError(
            "PDF extraction requires 'pdfplumber' or 'pypdf'. "
            "Please activate your virtual environment (e.g., .\\venv\\Scripts\\activate) "
            "or install dependencies via 'pip install -r requirements.txt'."
        )


def _read_docx(path: Path) -> str:
    import docx
    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.append(" ".join(c.text for c in row.cells))
    return "\n".join(parts)


def extract_text(path) -> str:
    """Read .pdf / .docx / .txt and return normalized text T.

    Scanned (image-only) PDFs yield empty text; OCR is out of scope for MTE."""
    path = Path(path)
    ext = path.suffix.lower()
    if ext == ".pdf":
        raw = _read_pdf(path)
    elif ext == ".docx":
        raw = _read_docx(path)
    elif ext in {".txt", ".md"}:
        raw = path.read_text(encoding="utf-8", errors="ignore")
    else:
        raise ValueError(f"Unsupported resume format: {ext}")
    text = normalize_text(raw)
    if not text:
        raise ValueError(f"No extractable text in {path.name} (scanned PDF?)")
    return text
